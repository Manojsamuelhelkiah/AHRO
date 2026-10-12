"""Convert the current portable learning guide to an editable Word document."""
from pathlib import Path
import sys
import re
from lxml import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(sys.argv[1] if len(sys.argv) > 1 else "deliverables/onboarding-control-tower-learning-guide.html")
OUTPUT = Path(sys.argv[2] if len(sys.argv) > 2 else "deliverables/onboarding-control-tower-learning-guide.docx")
root = html.parse(str(SOURCE))
doc = Document()
doc.core_properties.title = "Onboarding Control Tower — Product Learning Guide"
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    doc.styles[name].font.color.rgb = RGBColor.from_string("273B9D")


def page(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11.7 if landscape else 8.3)
    section.page_height = Inches(8.3 if landscape else 11.7)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)


page(doc.sections[0])
footer = doc.sections[0].footer.paragraphs[0]
footer.add_run("Onboarding Control Tower | User Guide | ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def text(value):
    return re.sub(r"\s+", " ", value or "")


def inline(node, paragraph, bold=False, italic=False):
    def add(value, b=bold, i=italic):
        if value:
            run = paragraph.add_run(text(value))
            run.bold, run.italic = b, i
    add(node.text)
    for child in node:
        if child.tag == "br":
            paragraph.add_run().add_break()
        else:
            inline(child, paragraph, bold or child.tag in ("strong", "b"),
                   italic or child.tag in ("em", "i"))
        add(child.tail)


def table(node):
    rows = node.xpath(".//tr")
    page(doc.add_section(WD_SECTION.NEW_PAGE), True)
    columns = max(len(row) for row in rows)
    grid = doc.add_table(rows=0, cols=columns)
    grid.style = "Table Grid"
    for row_index, row in enumerate(rows):
        cells = grid.add_row().cells
        for cell, source_cell in zip(cells, row):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            inline(source_cell, paragraph, row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(9)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            grid.rows[0]._tr.get_or_add_trPr().append(repeat)
    page(doc.add_section(WD_SECTION.NEW_PAGE))


def block(node):
    tag = node.tag
    if tag in ("h1", "h2", "h3", "h4"):
        level = 0 if tag == "h1" else int(tag[1]) - 1
        inline(node, doc.add_heading(level=level))
    elif tag == "p":
        inline(node, doc.add_paragraph())
    elif tag in ("ol", "ul"):
        for child in node:
            if child.tag == "li":
                inline(child, doc.add_paragraph(style="List Number" if tag == "ol" else "List Bullet"))
    elif tag == "table":
        table(node)
    else:
        for child in node:
            block(child)


block(root.find("body"))
doc.save(OUTPUT)

# Reopen the actual output and compare every text token to the HTML source.
verified = Document(OUTPUT)
from zipfile import ZipFile
from lxml import etree
with ZipFile(OUTPUT) as archive:
    xml = etree.fromstring(archive.read("word/document.xml"))
    word_text = " ".join(xml.xpath("//w:t/text()", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}))
source_text = " ".join(root.find("body").itertext())
assert source_text.split() == word_text.split(), "Word conversion lost or reordered text"
assert len(verified.tables) == len(root.xpath("//table"))
print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes); all text and both tables verified.")